import os
import sys
import json
import docx

# Reconfigure stdout to use UTF-8
sys.stdout.reconfigure(encoding='utf-8')

BASE_MAUDON_DIR = r'D:\legal\backend\pdf_maudon'
TARGET_TEMPLATE_DIR = r'D:\legal\backend\json_data\procedures'

os.makedirs(TARGET_TEMPLATE_DIR, exist_ok=True)

# Helper function to perform clean search and replace in a docx document
def replace_docx_placeholders(src_path, dst_path, replacements):
    doc = docx.Document(src_path)
    
    # Process regular paragraphs
    for p in doc.paragraphs:
        txt = p.text
        original = txt
        replaced = False
        for target, replacement in replacements:
            if target in txt:
                txt = txt.replace(target, replacement)
                replaced = True
        if replaced:
            p.text = txt
            
    # Process tables
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    txt = p.text
                    original = txt
                    replaced = False
                    for target, replacement in replacements:
                        if target in txt:
                            txt = txt.replace(target, replacement)
                            replaced = True
                    if replaced:
                        p.text = txt
                        
    doc.save(dst_path)
    print(f"  Successfully converted and saved {dst_path}")

# List of all template specifications
templates_specs = [
    # 1. khoi_kien_tranh_chap_lao_dong
    {
        "id": "khoi_kien_tranh_chap_lao_dong",
        "category": "lao_dong",
        "file_name": "Mẫu đơn khởi kiện tranh chấp lao động.docx",
        "title": "Đơn khởi kiện tranh chấp lao động",
        "description": "Đơn khởi kiện vụ án tranh chấp lao động (Mẫu số 23-DS) gửi Tòa án nhân dân giải quyết.",
        "replacements": [
            ("……(1), ngày….. tháng …… năm…….", "{{dia_diem_ngay}}"),
            ("Kính gửi: Toà án nhân dân (2)……………………………………", "Kính gửi: Toà án nhân dân {{toa_an}}"),
            ("Người khởi kiện: (3)", "Người khởi kiện: {{nguoi_khoi_kien}}"),
            ("Địa chỉ: (4)", "Địa chỉ: {{dia_chi_khoi_kien}}"),
            ("Số điện thoại: …………………(nếu có); số fax: ………………….(nếu có)", "Số điện thoại: {{sdt_khoi_kien}} (nếu có); số fax: (nếu có)"),
            ("Địa chỉ thư điện tử: ………................................................. (nếu có)", "Địa chỉ thư điện tử: {{email_khoi_kien}} (nếu có)"),
            ("Người bị kiện: (5)", "Người bị kiện: {{nguoi_bi_kien}}"),
            ("Địa chỉ (6)", "Địa chỉ: {{dia_chi_bi_kien}}"),
            ("Người có quyền, lợi ích được bảo vệ (nếu có)(7)", "Người có quyền, lợi ích được bảo vệ (nếu có): {{nguoi_duoc_bao_ve}}"),
            ("Địa chỉ: (8)", "Địa chỉ: {{dia_chi_duoc_bao_ve}}"),
            ("Người có quyền lợi, nghĩa vụ liên quan (nếu có) (9)", "Người có quyền lợi, nghĩa vụ liên quan (nếu có): {{nguoi_lien_quan}}"),
            ("Địa chỉ: (10)", "Địa chỉ: {{dia_chi_lien_quan}}"),
            ("Yêu cầu Tòa án giải quyết những vấn đề sau đây:(11)", "Yêu cầu Tòa án giải quyết những vấn đề sau đây: {{yeu_cau_toa_an}}"),
            ("Người làm chứng (nếu có) (12)", "Người làm chứng (nếu có): {{nguoi_lam_chung}}"),
            ("Địa chỉ: (13)", "Địa chỉ: {{dia_chi_lam_chung}}"),
            ("Danh mục tài liệu, chứng kèm theo đơn khởi kiện gồm có: (14)", "Danh mục tài liệu, chứng kèm theo đơn khởi kiện gồm có: {{danh_muc_tai_lieu}}"),
            ("(Các thông tin khác mà người khởi kiện xét thấy cần thiết cho việc giải quyết vụ án) (15)", "(Các thông tin khác mà người khởi kiện xét thấy cần thiết cho việc giải quyết vụ án): {{thong_tin_khac}}"),
        ],
        "fields": [
            {"key": "dia_diem_ngay", "label": "Địa điểm và ngày làm đơn", "question": "Nhập địa điểm và ngày tháng năm làm đơn (Ví dụ: Hà Nội, ngày 18 tháng 5 năm 2026):", "required": True},
            {"key": "toa_an", "label": "Tòa án có thẩm quyền", "question": "Kính gửi Tòa án nhân dân quận/huyện/tỉnh nào?", "required": True},
            {"key": "nguoi_khoi_kien", "label": "Họ tên người khởi kiện", "question": "Nhập Họ và tên của người khởi kiện (bạn):", "required": True},
            {"key": "dia_chi_khoi_kien", "label": "Địa chỉ người khởi kiện", "question": "Nhập địa chỉ cư trú đầy đủ của người khởi kiện:", "required": True},
            {"key": "sdt_khoi_kien", "label": "SĐT người khởi kiện", "question": "Nhập số điện thoại của người khởi kiện:", "required": True},
            {"key": "email_khoi_kien", "label": "Email người khởi kiện", "question": "Nhập địa chỉ email của người khởi kiện (nếu có):", "required": False},
            {"key": "nguoi_bi_kien", "label": "Họ tên người bị kiện", "question": "Nhập Họ và tên (hoặc tên Công ty/Đơn vị) của người bị kiện:", "required": True},
            {"key": "dia_chi_bi_kien", "label": "Địa chỉ người bị kiện", "question": "Nhập địa chỉ của người bị kiện (hoặc trụ sở công ty bị kiện):", "required": True},
            {"key": "nguoi_duoc_bao_ve", "label": "Người được bảo vệ (nếu có)", "question": "Nhập Họ tên người có quyền lợi ích được bảo vệ (nếu có):", "required": False},
            {"key": "dia_chi_duoc_bao_ve", "label": "Địa chỉ người được bảo vệ", "question": "Địa chỉ người được bảo vệ (nếu có):", "required": False},
            {"key": "nguoi_lien_quan", "label": "Người có quyền lợi nghĩa vụ liên quan", "question": "Họ tên người có quyền lợi, nghĩa vụ liên quan (nếu có):", "required": False},
            {"key": "dia_chi_lien_quan", "label": "Địa chỉ người liên quan", "question": "Địa chỉ người có quyền lợi nghĩa vụ liên quan (nếu có):", "required": False},
            {"key": "yeu_cau_toa_an", "label": "Nội dung yêu cầu giải quyết", "question": "Nhập các nội dung yêu cầu Tòa án giải quyết (Ví dụ: Buộc công ty thanh toán tiền lương còn nợ và bồi thường hợp đồng...):", "required": True},
            {"key": "nguoi_lam_chung", "label": "Người làm chứng (nếu có)", "question": "Họ tên người làm chứng (nếu có):", "required": False},
            {"key": "dia_chi_lam_chung", "label": "Địa chỉ người làm chứng", "question": "Địa chỉ người làm chứng (nếu có):", "required": False},
            {"key": "danh_muc_tai_lieu", "label": "Danh mục tài liệu chứng cứ", "question": "Nhập các tài liệu, chứng cứ kèm theo đơn (Ví dụ: Bản sao Hợp đồng lao động, Bảng thanh toán lương...):", "required": True},
            {"key": "thong_tin_khac", "label": "Các thông tin khác", "question": "Các thông tin khác xét thấy cần thiết (nếu có):", "required": False}
        ]
    },
    # 2. de_nghi_ki_tiep_hop_dong_lao_dong
    {
        "id": "de_nghi_ki_tiep_hop_dong_lao_dong",
        "category": "lao_dong",
        "file_name": "Mẫu đơn đề nghị kí tiếp hợp đồng lao động.docx",
        "title": "Đơn đề nghị ký tiếp hợp đồng lao động",
        "description": "Đơn đề xuất ban lãnh đạo công ty xem xét gia hạn, ký tiếp hợp đồng lao động mới khi hợp đồng cũ sắp hết hạn.",
        "replacements": [
            ("Kính gửi: Công ty................... ................", "Kính gửi: Công ty {{cong_ty}}"),
            ("Tôi tên là: ............................. Sinh năm: .....................................", "Tôi tên là: {{ho_ten}} Sinh năm: {{sinh_nam}}"),
            ("Hiện nay, tôi đang là nhân viên Công ty.............. ..............với chức danh/chức vụ là ................", "Hiện nay, tôi đang là nhân viên Công ty {{cong_ty}} với chức danh/chức vụ là {{chuc_vu}}"),
            ("Tôi đã ký hợp đồng lao động với Quý công ty từ ngày ............... đến ngày ..... Đến nay hợp đồng lao động đã sắp hết hạn.", "Tôi đã ký hợp đồng lao động với Quý công ty từ ngày {{tu_ngay}} đến ngày {{den_ngay}}. Đến nay hợp đồng lao động đã sắp hết hạn."),
            ("hoàn thành tốt công việc mà đã thoả thuận, ký kết trong hợp đồng lao động số...", "hoàn thành tốt công việc mà đã thoả thuận, ký kết trong hợp đồng lao động số {{so_hdld}}"),
            ("Những công việc đạt được ...........................", "Những công việc đạt được: {{cong_viec_dat_duoc}}"),
            ("để tôi được tiếp tục làm việc tại công ty với công việc ...............", "để tôi được tiếp tục làm việc tại công ty với công việc {{cong_viec_mong_muon}}"),
            ("..... Ngày .... tháng .... năm ..........", "{{ngay_lam_don}}"),
        ],
        "fields": [
            {"key": "cong_ty", "label": "Tên công ty", "question": "Nhập tên Công ty nơi bạn đang làm việc:", "required": True},
            {"key": "ho_ten", "label": "Họ và tên của bạn", "question": "Nhập Họ và tên đầy đủ của bạn:", "required": True},
            {"key": "sinh_nam", "label": "Năm sinh", "question": "Nhập năm sinh của bạn (Ví dụ: 1995):", "required": True},
            {"key": "chuc_vu", "label": "Chức vụ hiện tại", "question": "Nhập chức vụ/vị trí hiện tại của bạn:", "required": True},
            {"key": "tu_ngay", "label": "Hợp đồng từ ngày", "question": "Hợp đồng lao động hiện tại ký từ ngày nào (dd/mm/yyyy)?", "required": True},
            {"key": "den_ngay", "label": "Hợp đồng đến ngày", "question": "Hợp đồng lao động hiện tại hết hạn vào ngày nào (dd/mm/yyyy)?", "required": True},
            {"key": "so_hdld", "label": "Số hợp đồng lao động", "question": "Nhập số Hợp đồng lao động hiện tại của bạn:", "required": True},
            {"key": "cong_viec_dat_duoc", "label": "Thành tích, công việc đã đạt được", "question": "Nhập tóm tắt các kết quả, công việc nổi bật bạn đã đạt được tại công ty:", "required": True},
            {"key": "cong_viec_mong_muon", "label": "Công việc mong muốn tiếp tục làm", "question": "Nhập vị trí/công việc bạn muốn tiếp tục làm (Ví dụ: Nhân viên kỹ thuật):", "required": True},
            {"key": "ngay_lam_don", "label": "Ngày làm đơn", "question": "Nhập ngày làm đơn (Ví dụ: Hà Nội, ngày 18 tháng 5 năm 2026):", "required": True}
        ]
    },
    # 3. de_nghi_huong_tro_cap_that_nghiep
    {
        "id": "de_nghi_huong_tro_cap_that_nghiep",
        "category": "lao_dong",
        "file_name": "Đơn đề nghị hưởng trợ cấp thật nghiệp.docx",
        "title": "Đơn đề nghị hưởng trợ cấp thất nghiệp",
        "description": "Mẫu số 10 đề nghị hưởng trợ cấp thất nghiệp chính thức gửi đến Trung tâm Dịch vụ việc làm.",
        "replacements": [
            ("Kính gửi: Tổ chức Dịch vụ việc làm công …………", "Kính gửi: Tổ chức Dịch vụ việc làm công {{dich_vu_viec_lam}}"),
            ("Tên tôi là: …………………………………. Sinh ngày …/…./….", "Tên tôi là: {{ho_ten}} Sinh ngày {{ngay_sinh}}"),
            ("Số định danh cá nhân: ………………………….. Ngày cấp: …./…/….", "Số định danh cá nhân: {{so_cccd}} Ngày cấp: {{ngay_cap_cccd}}"),
            ("Số điện thoại: …………………….. Email (nếu có): ………………", "Số điện thoại: {{sdt}} Email (nếu có): {{email}}"),
            ("Trình độ đào tạo: …………………………………..……………", "Trình độ đào tạo: {{trinh_do_dao_tao}}"),
            ("Chuyên ngành đào tạo: …………………………………………", "Chuyên ngành đào tạo: {{chuyen_nganh_dao_tao}}"),
            ("Nơi thường trú: …………………………………………", "Nơi thường trú: {{noi_thuong_tru}}"),
            ("Nơi ở hiện nay: …………………………………………………", "Nơi ở hiện nay: {{noi_o_hien_nay}}"),
            ("Ngày …/…/…. tôi đã chấm dứt hợp đồng lao động/hợp đồng làm việc với (tên đơn vị) …………… tại địa chỉ: ……………………………", "Ngày {{ngay_cham_dut_hdld}} tôi đã chấm dứt hợp đồng lao động/hợp đồng làm việc với (tên đơn vị) {{ten_don_vi}} tại địa chỉ: {{dia_chi_don_vi}}"),
            ("Lý do chấm dứt hợp đồng lao động/hợp đồng làm việc: ………………", "Lý do chấm dứt hợp đồng lao động/hợp đồng làm việc: {{ly_do_cham_dut}}"),
            ("Loại hợp đồng lao động/hợp đồng làm việc: ……………………….", "Loại hợp đồng lao động/hợp đồng làm việc: {{loai_hdld}}"),
            ("Số tháng đóng bảo hiểm thất nghiệp: ... tháng.", "Số tháng đóng bảo hiểm thất nghiệp: {{so_thang_dong_bhtn}} tháng."),
            ("- Tên chủ tài khoản: ……………………….. Số tài khoản: ……………", "- Tên chủ tài khoản: {{ten_chu_tk}} Số tài khoản: {{so_tk}}"),
            ("- Tại ngân hàng: ……………………………………………….", "- Tại ngân hàng: {{ngan_hang}}"),
            ("Trường hợp không có tài khoản, nhận tiền mặt tại BHXH cấp cơ sở: …………………..", "Trường hợp không có tài khoản, nhận tiền mặt tại BHXH cấp cơ sở: {{nhan_tien_mat_tai}}"),
            ("Nơi đăng ký khám chữa bệnh ban đầu: …………………………", "Nơi đăng ký khám chữa bệnh ban đầu: {{noi_kham_chua_benh_ban_dau}}"),
            ("..., ngày ... tháng ... năm ...", "{{ngay_lam_don}}")
        ],
        "fields": [
            {"key": "dich_vu_viec_lam", "label": "Trung tâm Dịch vụ việc làm", "question": "Kính gửi Trung tâm Dịch vụ việc làm tỉnh/thành phố nào?", "required": True},
            {"key": "ho_ten", "label": "Họ và tên người lao động", "question": "Nhập Họ và tên đầy đủ của bạn:", "required": True},
            {"key": "ngay_sinh", "label": "Ngày sinh", "question": "Nhập ngày tháng năm sinh của bạn (dd/mm/yyyy):", "required": True},
            {"key": "so_cccd", "label": "Số định danh / CCCD", "question": "Nhập số CCCD/định danh cá nhân của bạn:", "required": True},
            {"key": "ngay_cap_cccd", "label": "Ngày cấp CCCD", "question": "CCCD được cấp vào ngày nào (dd/mm/yyyy)?", "required": True},
            {"key": "sdt", "label": "Số điện thoại", "question": "Nhập số điện thoại liên lạc của bạn:", "required": True},
            {"key": "email", "label": "Email", "question": "Nhập email của bạn (nếu có):", "required": False},
            {"key": "trinh_do_dao_tao", "label": "Trình độ đào tạo", "question": "Trình độ đào tạo của bạn (Ví dụ: Đại học, Cao đẳng, Phổ thông...):", "required": True},
            {"key": "chuyen_nganh_dao_tao", "label": "Chuyên ngành đào tạo", "question": "Chuyên ngành đào tạo của bạn (Ví dụ: Công nghệ thông tin...):", "required": True},
            {"key": "noi_thuong_tru", "label": "Nơi thường trú", "question": "Nhập địa chỉ nơi thường trú theo hộ khẩu/CCCD:", "required": True},
            {"key": "noi_o_hien_nay", "label": "Nơi ở hiện nay", "question": "Nhập địa chỉ nơi cư trú hiện tại của bạn:", "required": True},
            {"key": "ngay_cham_dut_hdld", "label": "Ngày chấm dứt hợp đồng", "question": "Ngày chấm dứt hợp đồng lao động là ngày nào (dd/mm/yyyy)?", "required": True},
            {"key": "ten_don_vi", "label": "Tên công ty cũ", "question": "Tên công ty/đơn vị cũ nơi bạn vừa thôi việc:", "required": True},
            {"key": "dia_chi_don_vi", "label": "Địa chỉ công ty cũ", "question": "Nhập địa chỉ trụ sở của công ty/đơn vị cũ:", "required": True},
            {"key": "ly_do_cham_dut", "label": "Lý do chấm dứt hợp đồng", "question": "Nhập lý do chấm dứt hợp đồng (Ví dụ: Hết hạn hợp đồng, Đơn phương xin nghỉ...):", "required": True},
            {"key": "loai_hdld", "label": "Loại hợp đồng lao động", "question": "Loại hợp đồng lao động của bạn (Ví dụ: Không xác định thời hạn, 12 tháng...):", "required": True},
            {"key": "so_thang_dong_bhtn", "label": "Số tháng đóng BHTN", "question": "Tổng số tháng bạn đóng Bảo hiểm thất nghiệp là bao nhiêu?", "required": True},
            {"key": "ten_chu_tk", "label": "Tên chủ tài khoản ngân hàng", "question": "Nhập tên chủ tài khoản nhận trợ cấp (Ví dụ: NGUYEN VAN A):", "required": True},
            {"key": "so_tk", "label": "Số tài khoản ngân hàng", "question": "Nhập số tài khoản ngân hàng:", "required": True},
            {"key": "ngan_hang", "label": "Tên ngân hàng", "question": "Nhập tên ngân hàng (Ví dụ: Vietcombank chi nhánh Hà Nội):", "required": True},
            {"key": "nhan_tien_mat_tai", "label": "Nơi nhận tiền mặt", "question": "Nếu không nhận qua thẻ, bạn muốn nhận tiền mặt tại BHXH quận/huyện nào?", "required": False},
            {"key": "noi_kham_chua_benh_ban_dau", "label": "Nơi đăng ký KCB ban đầu", "question": "Nhập nơi đăng ký khám chữa bệnh ban đầu ghi trên thẻ BHYT của bạn:", "required": True},
            {"key": "ngay_lam_don", "label": "Ngày làm đơn", "question": "Nhập ngày làm đơn (Ví dụ: Hà Nội, ngày 18 tháng 5 năm 2026):", "required": True}
        ]
    },
    # 4. don_khieu_nai
    {
        "id": "don_khieu_nai",
        "category": "dan_su",
        "file_name": "Mau-don-khieu-nai.docx",
        "title": "Đơn khiếu nại (Mẫu số 01)",
        "description": "Mẫu đơn khiếu nại hành chính chính thức (Mẫu số 01) dùng để khiếu nại quyết định hoặc hành vi hành chính.",
        "replacements": [
            ("..., ngày... tháng... năm...", "{{ngay_lam_don}}"),
            ("Kính gửi: ………………….. (1)", "Kính gửi: {{kinh_gui}}"),
            ("- Họ và tên .................................................................................................…", "- Họ và tên: {{ho_ten_nguoi_khieu_nai}}"),
            ("- Địa chỉ:.........................................................................................................", "- Địa chỉ: {{dia_chi_nguoi_khieu_nai}}"),
            ("- Số căn cước/định danh cá nhân/giấy chứng nhận căn cước/hộ chiếu/....... , ngày cấp............., nơi cấp:………………………………..", "- Số căn cước/định danh cá nhân: {{cccd_nguoi_khieu_nai}}, ngày cấp: {{ngay_cap_cccd_nguoi_khieu_nai}}, nơi cấp: {{noi_cap_cccd_nguoi_khieu_nai}}"),
            ("- Họ và tên:………………………..…………………………………………", "- Họ và tên: {{ho_ten_nguoi_uy_quyen}}"),
            ("- Địa chỉ:.........................................................................................................;", "- Địa chỉ: {{dia_chi_nguoi_uy_quyen}};"),
            ("- Số căn cước/định danh cá nhân/giấy chứng nhận căn cước/hộ chiếu/........, ngày cấp............., nơi cấp:………………………………..", "- Số căn cước/định danh cá nhân: {{cccd_nguoi_uy_quyen}}, ngày cấp: {{ngay_cap_cccd_nguoi_uy_quyen}}, nơi cấp: {{noi_cap_cccd_nguoi_uy_quyen}}"),
            ("- Tên cơ quan, tổ chức, cá nhân:……………………………………………", "- Tên cơ quan, tổ chức, cá nhân: {{ten_bi_khieu_nai}}"),
            ("- Địa chỉ:.........................................................................................................", "- Địa chỉ: {{dia_chi_bi_khieu_nai}}"),
            ("4. Khiếu nại về việc:................................................................................. (5)", "4. Khiếu nại về việc: {{khieu_nai_ve_viec}}"),
            ("5. Nội dung khiếu nại:.............................................................................. (6)", "5. Nội dung khiếu nại: {{noi_dung_khieu_nai}}")
        ],
        "fields": [
            {"key": "ngay_lam_don", "label": "Ngày làm đơn", "question": "Nhập địa điểm và ngày làm đơn (Ví dụ: Hà Nội, ngày 18 tháng 5 năm 2026):", "required": True},
            {"key": "kinh_gui", "label": "Cơ quan giải quyết khiếu nại", "question": "Nhập tên cơ quan, tổ chức, cá nhân có thẩm quyền giải quyết khiếu nại (Kính gửi ai?):", "required": True},
            {"key": "ho_ten_nguoi_khieu_nai", "label": "Họ tên người khiếu nại", "question": "Nhập Họ và tên đầy đủ của người khiếu nại:", "required": True},
            {"key": "dia_chi_nguoi_khieu_nai", "label": "Địa chỉ người khiếu nại", "question": "Nhập địa chỉ của người khiếu nại:", "required": True},
            {"key": "cccd_nguoi_khieu_nai", "label": "CCCD người khiếu nại", "question": "Nhập số CCCD/Định danh cá nhân của người khiếu nại:", "required": True},
            {"key": "ngay_cap_cccd_nguoi_khieu_nai", "label": "Ngày cấp CCCD", "question": "Nhập ngày cấp CCCD (dd/mm/yyyy):", "required": True},
            {"key": "noi_cap_cccd_nguoi_khieu_nai", "label": "Nơi cấp CCCD", "question": "CCCD do đơn vị nào cấp (Ví dụ: Cục CSQLHC về trật tự xã hội):", "required": True},
            {"key": "ho_ten_nguoi_uy_quyen", "label": "Họ tên người được ủy quyền (nếu có)", "question": "Nhập Họ tên người được ủy quyền khiếu nại (bỏ qua nếu tự khiếu nại):", "required": False},
            {"key": "dia_chi_nguoi_uy_quyen", "label": "Địa chỉ người ủy quyền (nếu có)", "question": "Địa chỉ người được ủy quyền (nếu có):", "required": False},
            {"key": "cccd_nguoi_uy_quyen", "label": "CCCD người ủy quyền (nếu có)", "question": "Số CCCD người được ủy quyền (nếu có):", "required": False},
            {"key": "ngay_cap_cccd_nguoi_uy_quyen", "label": "Ngày cấp CCCD người ủy quyền", "question": "Ngày cấp CCCD người ủy quyền (nếu có):", "required": False},
            {"key": "noi_cap_cccd_nguoi_uy_quyen", "label": "Nơi cấp CCCD người ủy quyền", "question": "Nơi cấp CCCD người ủy quyền (nếu có):", "required": False},
            {"key": "ten_bi_khieu_nai", "label": "Bên bị khiếu nại", "question": "Nhập tên của cơ quan, tổ chức hoặc cá nhân bị khiếu nại:", "required": True},
            {"key": "dia_chi_bi_khieu_nai", "label": "Địa chỉ bên bị khiếu nại", "question": "Nhập địa chỉ của bên bị khiếu nại:", "required": True},
            {"key": "khieu_nai_ve_viec", "label": "Khiếu nại về việc", "question": "Nội dung khiếu nại về quyết định/hành vi hành chính gì? (Ví dụ: Quyết định thu hồi đất số 123/QĐ-UBND):", "required": True},
            {"key": "noi_dung_khieu_nai", "label": "Nội dung khiếu nại", "question": "Trình bày chi tiết lý do bạn khiếu nại và các yêu cầu cụ thể giải quyết khiếu nại:", "required": True}
        ]
    },
    # 5. don_xin_tam_hoan_nvqs
    {
        "id": "don_xin_tam_hoan_nvqs",
        "category": "dan_su",
        "file_name": "don-xin-tam-hoan-nvqs.docx",
        "title": "Đơn xin tạm hoãn nghĩa vụ quân sự",
        "description": "Đơn đề nghị Hội đồng nghĩa vụ quân sự xem xét tạm hoãn gọi nhập ngũ theo Luật Nghĩa vụ quân sự 2015.",
        "replacements": [
            ("Kính gửi: 	- Hội đồng Nghĩa vụ quân sự quận/huyện ……….;", "Kính gửi: \t- Hội đồng Nghĩa vụ quân sự quận/huyện {{huyen_quan}};"),
            ("- Hội đồng Nghĩa vụ quân sự xã/phường………...", "- Hội đồng Nghĩa vụ quân sự xã/phường {{xa_phuong}}."),
            ("Tên tôi là: ................................................... Ngày, tháng, năm sinh: .........................", "Tên tôi là: {{ho_ten}} Ngày, tháng, năm sinh: {{ngay_sinh}}"),
            ("CMTND/CCCD số: .............................. ngày cấp: ............. nơi cấp: .........................", "CMTND/CCCD số: {{so_cccd}} ngày cấp: {{ngay_cap_cccd}} nơi cấp: {{noi_cap_cccd}}"),
            ("Số điện thoại: ......................................................................................................", "Số điện thoại: {{sdt}}"),
            ("Quê quán: ....................................................................................................................", "Quê quán: {{que_quan}}"),
            ("Địa chỉ thường trú: ....................................................................................................", "Địa chỉ thường trú: {{noi_thuong_tru}}"),
            ("Nơi ở hiện nay: ...........................................................................................................", "Nơi ở hiện nay: {{noi_o_hien_nay}}"),
            ("Nay tôi làm đơn này kính mong Hội đồng NVQS xã/phường…….xem xét cho tôi được tạm hoãn, miễn gọi nhập ngũ năm…….", "Nay tôi làm đơn này kính mong Hội đồng NVQS xã/phường {{xa_phuong}} xem xét cho tôi được tạm hoãn, miễn gọi nhập ngũ năm {{nam_nhap_ngu}}"),
            ("Lý do: ......................................................................................................................... .....................................................................................................................................theo Điều 41 Luật nghĩa vụ quân sự 2015.", "Lý do: {{ly_do}} theo Điều 41 Luật nghĩa vụ quân sự 2015."),
            ("Hiện tại đang theo học tại Trường: ..................................................................", "Hiện tại đang theo học tại Trường: {{truong_hoc}}"),
            ("Mã số sinh viên: ........................ Lớp: ........................ Khoa: .............................", "Mã số sinh viên: {{ms_sinh_vien}} Lớp: {{lop}} Khoa: {{khoa}}"),
            (".............................................. Thời gian đào tạo: từ 20.... đến 20…", "Thời gian đào tạo: từ 20{{nam_bat_dau_dao_tao}} đến 20{{nam_ket_thuc_dao_tao}}"),
            ("..............., ngày ..... tháng ..... năm .......", "{{ngay_lam_don}}")
        ],
        "fields": [
            {"key": "huyen_quan", "label": "Huyện/Quận", "question": "Nhập tên Quận/Huyện của Hội đồng NVQS tiếp nhận đơn:", "required": True},
            {"key": "xa_phuong", "label": "Xã/Phường", "question": "Nhập tên Xã/Phường của Hội đồng NVQS tiếp nhận đơn:", "required": True},
            {"key": "ho_ten", "label": "Họ và tên của bạn", "question": "Nhập Họ và tên đầy đủ của bạn:", "required": True},
            {"key": "ngay_sinh", "label": "Ngày sinh", "question": "Nhập ngày tháng năm sinh của bạn (dd/mm/yyyy):", "required": True},
            {"key": "so_cccd", "label": "Số CCCD", "question": "Nhập số CCCD của bạn:", "required": True},
            {"key": "ngay_cap_cccd", "label": "Ngày cấp CCCD", "question": "Nhập ngày cấp CCCD (dd/mm/yyyy):", "required": True},
            {"key": "noi_cap_cccd", "label": "Nơi cấp CCCD", "question": "Nhập nơi cấp CCCD (Ví dụ: Cục Cảnh sát QLHC về trật tự xã hội):", "required": True},
            {"key": "sdt", "label": "Số điện thoại", "question": "Nhập số điện thoại của bạn:", "required": True},
            {"key": "que_quan", "label": "Quê quán", "question": "Nhập quê quán của bạn (Tỉnh/Thành phố):", "required": True},
            {"key": "noi_thuong_tru", "label": "Nơi thường trú", "question": "Nhập địa chỉ nơi đăng ký thường trú của bạn:", "required": True},
            {"key": "noi_o_hien_nay", "label": "Nơi ở hiện nay", "question": "Nhập địa chỉ nơi ở hiện nay của bạn:", "required": True},
            {"key": "nam_nhap_ngu", "label": "Năm gọi nhập ngũ", "question": "Nhập năm tuyển quân bạn muốn hoãn gọi nhập ngũ (Ví dụ: 2026):", "required": True},
            {"key": "ly_do", "label": "Lý do hoãn nghĩa vụ", "question": "Lý do xin tạm hoãn của bạn (Ví dụ: Đang là sinh viên theo học hệ đào tạo chính quy):", "required": True},
            {"key": "truong_hoc", "label": "Tên trường đại học/cao đẳng", "question": "Nhập tên trường bạn đang theo học (nếu hoãn do đi học):", "required": False},
            {"key": "ms_sinh_vien", "label": "Mã số sinh viên", "question": "Nhập mã số sinh viên của bạn (nếu có):", "required": False},
            {"key": "lop", "label": "Lớp học", "question": "Nhập tên lớp học hiện tại của bạn:", "required": False},
            {"key": "khoa", "label": "Khoa", "question": "Nhập tên khoa đào tạo của bạn:", "required": False},
            {"key": "nam_bat_dau_dao_tao", "label": "Năm bắt đầu học", "question": "Hai số cuối của năm bắt đầu khóa học (Ví dụ: 22 cho năm 2022):", "required": False},
            {"key": "nam_ket_thuc_dao_tao", "label": "Năm kết thúc học", "question": "Hai số cuối của năm kết thúc khóa học (Ví dụ: 26 cho năm 2026):", "required": False},
            {"key": "ngay_lam_don", "label": "Ngày làm đơn", "question": "Nhập địa điểm và ngày làm đơn (Ví dụ: Hà Nội, ngày 18 tháng 5 năm 2026):", "required": True}
        ]
    },
    # 6. mau_don_to_cao
    {
        "id": "mau_don_to_cao",
        "category": "dan_su",
        "file_name": "mau-don-to-cao-1.docx",
        "title": "Đơn tố cáo hành vi lừa đảo",
        "description": "Đơn tố cáo gửi cơ quan Công an có thẩm quyền về hành vi lừa đảo chiếm đoạt tài sản.",
        "replacements": [
            ("….., ngày......, tháng......, năm 20.....", "{{ngay_lam_don}}"),
            ("(Về hành vi lừa đảo chiếm đoạt tài sản của ông/bà ...........)", "(Về hành vi lừa đảo chiếm đoạt tài sản của ông/bà {{ten_nguoi_bi_to_cao_tieu_de}})"),
            ("Kính gửi: Công an tỉnh/ thành phố …………………", "Kính gửi: Công an tỉnh/ thành phố {{cong_an_tinh_tp}}"),
            ("Họ và tên: ...............................................................", "Họ và tên: {{ho_ten_nguoi_to_cao}}"),
            ("Ngày sinh: ...............................................................", "Ngày sinh: {{ngay_sinh_nguoi_to_cao}}"),
            ("CCCD/CC số: .....................................................", "CCCD/CC số: {{cccd_nguoi_to_cao}}"),
            ("Địa chỉ thường trú: ................................................", "Địa chỉ thường trú: {{dia_chi_thuong_tru_nguoi_to_cao}}"),
            ("Số điện thoại: .........................................................", "Số điện thoại: {{sdt_nguoi_to_cao}}"),
            ("Họ và tên: ...............................................................", "Họ và tên: {{ho_ten_nguoi_bi_to_cao}}"),
            ("Ngày sinh (nếu biết): ............................................", "Ngày sinh (nếu biết): {{ngay_sinh_nguoi_bi_to_cao}}"),
            ("Địa chỉ: ....................................................................", "Địa chỉ: {{dia_chi_nguoi_bi_to_cao}}"),
            ("Thông tin khác (số điện thoại, tài khoản, mạng xã hội…): ....................................................", "Thông tin khác (số điện thoại, tài khoản, mạng xã hội…): {{thong_tin_khac_nguoi_bi_to_cao}}"),
            ("Tôi làm đơn này tố cáo hành vi có dấu hiệu lừa đảo chiếm đoạt tài sản của ông/bà ........................................ với nội dung như sau:", "Tôi làm đơn này tố cáo hành vi có dấu hiệu lừa đảo chiếm đoạt tài sản của ông/bà {{ho_ten_nguoi_bi_to_cao_2}} với nội dung như sau:"),
            ("Vào khoảng thời gian từ ngày …… đến ngày ……, ông/bà …………… đã sử dụng các thủ đoạn gian dối như:", "Vào khoảng thời gian từ ngày {{ngay_dau}} đến ngày {{ngay_cuoi}}, ông/bà {{ho_ten_nguoi_bi_to_cao_3}} đã sử dụng các thủ đoạn gian dối như: {{hanh_vi_gian_doi}}"),
            ("Số tiền: ....................................................", "Số tiền: {{so_tien_thiet_hai}}"),
            ("Hình thức chuyển (chuyển khoản/tiền mặt…): ................", "Hình thức chuyển (chuyển khoản/tiền mặt…): {{hinh_thuc_chuyen}}"),
            ("Tổng giá trị thiệt hại: .....................................", "Tổng giá trị thiệt hại: {{tong_gia_tri_thiet_hai}}"),
            ("Sau khi nhận được tài sản, ông/bà …………… đã:", "Sau khi nhận được tài sản, ông/bà {{ho_ten_nguoi_bi_to_cao_4}} đã:"),
            ("- Tài liệu khác: ................................................", "- Tài liệu khác: {{tai_lieu_khac}}"),
        ],
        "fields": [
            {"key": "ngay_lam_don", "label": "Ngày làm đơn", "question": "Nhập địa điểm và ngày làm đơn (Ví dụ: Hà Nội, ngày 18 tháng 5 năm 2026):", "required": True},
            {"key": "ten_nguoi_bi_to_cao_tieu_de", "label": "Tên người bị tố cáo (tiêu đề)", "question": "Nhập Họ tên của người bị tố cáo (Ví dụ: Nguyễn Văn B):", "required": True},
            {"key": "cong_an_tinh_tp", "label": "Cơ quan Công an tiếp nhận", "question": "Kính gửi Công an Tỉnh/Thành phố nào?", "required": True},
            {"key": "ho_ten_nguoi_to_cao", "label": "Họ tên người tố cáo", "question": "Họ và tên đầy đủ của bạn (người tố cáo):", "required": True},
            {"key": "ngay_sinh_nguoi_to_cao", "label": "Ngày sinh người tố cáo", "question": "Ngày sinh của bạn (dd/mm/yyyy):", "required": True},
            {"key": "cccd_nguoi_to_cao", "label": "CCCD người tố cáo", "question": "Số CCCD/Định danh cá nhân của bạn:", "required": True},
            {"key": "dia_chi_thuong_tru_nguoi_to_cao", "label": "Địa chỉ người tố cáo", "question": "Địa chỉ thường trú đầy đủ của bạn:", "required": True},
            {"key": "sdt_nguoi_to_cao", "label": "Số điện thoại của bạn", "question": "Số điện thoại liên lạc của bạn:", "required": True},
            {"key": "ho_ten_nguoi_bi_to_cao", "label": "Họ tên người bị tố cáo", "question": "Nhập Họ và tên đầy đủ của người bị tố cáo:", "required": True},
            {"key": "ngay_sinh_nguoi_bi_to_cao", "label": "Ngày sinh người bị tố cáo", "question": "Ngày sinh người bị tố cáo (nếu biết, nếu không ghi 'Không rõ'):", "required": False},
            {"key": "dia_chi_nguoi_bi_to_cao", "label": "Địa chỉ người bị tố cáo", "question": "Địa chỉ cư trú hoặc trụ sở làm việc của người bị tố cáo (nếu biết):", "required": True},
            {"key": "thong_tin_khac_nguoi_bi_to_cao", "label": "Thông tin liên lạc khác", "question": "Số điện thoại, tài khoản ngân hàng, link Facebook... của người bị tố cáo:", "required": False},
            {"key": "ho_ten_nguoi_bi_to_cao_2", "label": "Xác nhận tên người bị tố cáo", "question": "Xác nhận lại tên người bị tố cáo một lần nữa:", "required": True},
            {"key": "ngay_dau", "label": "Bắt đầu từ ngày", "question": "Hành vi lừa đảo diễn ra từ ngày nào (dd/mm/yyyy)?", "required": True},
            {"key": "ngay_cuoi", "label": "Đến ngày", "question": "Đến ngày nào (dd/mm/yyyy)?", "required": True},
            {"key": "ho_ten_nguoi_bi_to_cao_3", "label": "Tên người bị tố cáo (hành vi)", "question": "Xác nhận lại tên người bị tố cáo (cho mục diễn giải hành vi):", "required": True},
            {"key": "hanh_vi_gian_doi", "label": "Chi tiết hành vi gian dối", "question": "Mô tả chi tiết các thủ đoạn, hành vi lừa đảo của đối tượng (Ví dụ: hứa hẹn xin việc, bán đất khống...):", "required": True},
            {"key": "so_tien_thiet_hai", "label": "Số tiền bị lừa", "question": "Nhập số tiền hoặc mô tả tài sản bạn đã chuyển giao bị đối tượng chiếm đoạt:", "required": True},
            {"key": "hinh_thuc_chuyen", "label": "Hình thức chuyển tiền", "question": "Hình thức chuyển giao tài sản (Ví dụ: chuyển khoản ngân hàng, tiền mặt...):", "required": True},
            {"key": "tong_gia_tri_thiet_hai", "label": "Tổng giá trị thiệt hại", "question": "Tổng giá trị thiệt hại bằng tiền (Ví dụ: 200.000.000 đồng):", "required": True},
            {"key": "ho_ten_nguoi_bi_to_cao_4", "label": "Tên người bị tố cáo (sau nhận tiền)", "question": "Xác nhận lại tên người bị tố cáo (cho mục sau khi nhận tài sản):", "required": True},
            {"key": "tai_lieu_khac", "label": "Tài liệu chứng cứ khác", "question": "Nhập tên các tài liệu, chứng cứ kèm theo đơn (sao kê tài khoản, tin nhắn...):", "required": True}
        ]
    },
    # 7. don_to_giac_toi_pham
    {
        "id": "don_to_giac_toi_pham",
        "category": "dan_su",
        "file_name": "ĐƠN TỐ GIÁC TỘI PHẠM.docx",
        "title": "Đơn tố giác tội phạm",
        "description": "Mẫu đơn trình báo, tố giác tội phạm gửi cơ quan Cảnh sát điều tra, Công an cấp xã/phường/quận.",
        "replacements": [
            ("……….., ngày ….. tháng …. năm…..", "{{ngay_lam_don}}"),
            ("(Về hành vi……………………) (1)", "(Về hành vi {{hanh_vi_vi_pham}})"),
            ("Kính gửi: Cơ quan điều tra, Công an xã/phường (2) …………...", "Kính gửi: Cơ quan điều tra, Công an xã/phường {{cong_an_xa_phuong}}"),
            ("Tôi tên là:…………………………………Sinh năm:………………………………", "Tôi tên là: {{ho_ten_nguoi_lam_don}} Sinh năm: {{sinh_nam_nguoi_lam_don}}"),
            ("CCCD số:……………………..do:………………..cấy ngày:…………………...…", "CCCD số: {{cccd_nguoi_lam_don}} do: {{cap_boi_nguoi_lam_don}} cấp ngày: {{ngay_cap_cccd_nguoi_lam_don}}"),
            ("Nơi thường trú:………………………………………………….………………", "Nơi thường trú: {{noi_thuong_tru_nguoi_lam_don}}"),
            ("Hiện đang cư ngụ tại:…………………………………………………………………", "Hiện đang cư ngụ tại: {{tam_tru_nguoi_lam_don}}"),
            ("Họ và tên:…………………….....", "Họ và tên đối tượng: {{ho_ten_doi_tuong}}"),
            ("Hiện đang cư ngụ tại:..………………………………………………………………", "Hiện cư ngụ tại: {{dia_chi_doi_tuong}}"),
            ("Đối tượng này đã có hành vi (3)……………………...………………………………", "Đối tượng này đã có hành vi: {{chi_tiet_hanh_vi}}"),
            ("Chứng cứ chứng minh (nếu có) (4):…………………….……………….……………", "Chứng cứ chứng minh (nếu có): {{chung_cu}}")
        ],
        "fields": [
            {"key": "ngay_lam_don", "label": "Ngày làm đơn", "question": "Nhập địa điểm và ngày làm đơn (Ví dụ: Hà Nội, ngày 18 tháng 5 năm 2026):", "required": True},
            {"key": "hanh_vi_vi_pham", "label": "Hành vi vi phạm", "question": "Nhập hành vi phạm tội cần tố giác (Ví dụ: cướp giật tài sản, cố ý gây thương tích...):", "required": True},
            {"key": "cong_an_xa_phuong", "label": "Cơ quan công an nhận đơn", "question": "Kính gửi Công an xã/phường/thị trấn nào?", "required": True},
            {"key": "ho_ten_nguoi_lam_don", "label": "Họ tên người làm đơn", "question": "Nhập Họ và tên đầy đủ của bạn (người tố giác):", "required": True},
            {"key": "sinh_nam_nguoi_lam_don", "label": "Năm sinh của bạn", "question": "Nhập năm sinh của bạn (Ví dụ: 1993):", "required": True},
            {"key": "cccd_nguoi_lam_don", "label": "CCCD của bạn", "question": "Nhập số CCCD của bạn:", "required": True},
            {"key": "cap_boi_nguoi_lam_don", "label": "Nơi cấp CCCD", "question": "CCCD của bạn do cơ quan nào cấp?", "required": True},
            {"key": "ngay_cap_cccd_nguoi_lam_don", "label": "Ngày cấp CCCD", "question": "Nhập ngày cấp CCCD của bạn (dd/mm/yyyy):", "required": True},
            {"key": "noi_thuong_tru_nguoi_lam_don", "label": "Nơi thường trú của bạn", "question": "Nhập địa chỉ đăng ký thường trú của bạn:", "required": True},
            {"key": "tam_tru_nguoi_lam_don", "label": "Địa chỉ cư trú hiện tại", "question": "Địa chỉ nơi bạn đang sinh sống, cư ngụ hiện nay:", "required": True},
            {"key": "ho_ten_doi_tuong", "label": "Tên đối tượng bị tố giác", "question": "Nhập Họ và tên đối tượng vi phạm (nếu không rõ, ghi 'Chưa rõ nhân thân'):", "required": True},
            {"key": "dia_chi_doi_tuong", "label": "Địa chỉ đối tượng", "question": "Nhập địa chỉ cư trú của đối tượng vi phạm (nếu biết, nếu không ghi 'Không rõ'):", "required": True},
            {"key": "chi_tiet_hanh_vi", "label": "Mô tả chi tiết hành vi phạm tội", "question": "Mô tả chi tiết diễn biến vụ việc và hành vi vi phạm pháp luật của đối tượng:", "required": True},
            {"key": "chung_cu", "label": "Chứng cứ kèm theo", "question": "Nhập các tài liệu, chứng cứ kèm theo để chứng minh (nếu có):", "required": False}
        ]
    },
    # 8. de_nghi_giai_quyet_tai_nan_giao_thong
    {
        "id": "de_nghi_giai_quyet_tai_nan_giao_thong",
        "category": "giao_thong",
        "file_name": "MAU DON DE NGHI GIAI QUYET TAI NAN GIAO THONG.docx",
        "title": "Đơn đề nghị giải quyết tai nạn giao thông",
        "description": "Đơn trình báo và yêu cầu cơ quan Công an tiến hành điều tra, giải quyết vụ tai nạn giao thông xảy ra đối với bạn.",
        "replacements": [
            ("… , ngày … tháng … năm …", "{{ngay_lam_don}}"),
            ("Kính gửi:\t", "Kính gửi: {{kinh_gui}}"),
            ("Họ và tên người đề nghị:\t", "Họ và tên người đề nghị: {{ho_ten_nguoi_de_nghi}}"),
            ("Sinh năm:\t", "Sinh năm: {{sinh_nam}}"),
            ("CMND số: \tNgày cấp: \tNơi cấp: \t", "CMND/CCCD số: {{cmnd_so}} Ngày cấp: {{ngay_cap}} Nơi cấp: {{noi_cap}}"),
            ("Địa chỉ cư trú: \t", "Địa chỉ cư trú: {{dia_chi_cu_tru}}"),
            ("Số điện thoại:\t", "Số điện thoại: {{sdt}}"),
            ("Trình bày nội dung sự việc", "Trình bày nội dung sự việc: {{noi_dung_su_viec}}"),
            ("đối với \t", "đối với {{nguoi_gay_tai_nan}} (người gây ra tai nạn giao thông)"),
            ("\t( người gây ra tai nạn giao thông) về hành vi\t\t trước pháp luật. ", "về hành vi {{hanh_vi_vi_pham}} trước pháp luật.")
        ],
        "fields": [
            {"key": "ngay_lam_don", "label": "Ngày làm đơn", "question": "Nhập địa điểm và ngày làm đơn (Ví dụ: Hà Nội, ngày 18 tháng 5 năm 2026):", "required": True},
            {"key": "kinh_gui", "label": "Cơ quan công an giải quyết", "question": "Kính gửi Đội Cảnh sát giao thông / Công an quận/huyện nào?", "required": True},
            {"key": "ho_ten_nguoi_de_nghi", "label": "Họ tên người đề nghị", "question": "Nhập Họ và tên đầy đủ của bạn (người đề nghị giải quyết):", "required": True},
            {"key": "sinh_nam", "label": "Năm sinh", "question": "Nhập năm sinh của bạn:", "required": True},
            {"key": "cmnd_so", "label": "Số CMND/CCCD của bạn", "question": "Nhập số CMND/CCCD của bạn:", "required": True},
            {"key": "ngay_cap", "label": "Ngày cấp", "question": "Nhập ngày cấp CMND/CCCD (dd/mm/yyyy):", "required": True},
            {"key": "noi_cap", "label": "Nơi cấp", "question": "Nhập nơi cấp CMND/CCCD (Ví dụ: Cục Cảnh sát QLHC về trật tự xã hội):", "required": True},
            {"key": "dia_chi_cu_tru", "label": "Địa chỉ cư trú", "question": "Nhập địa chỉ cư trú đầy đủ hiện tại của bạn:", "required": True},
            {"key": "sdt", "label": "Số điện thoại", "question": "Nhập số điện thoại liên lạc của bạn:", "required": True},
            {"key": "noi_dung_su_viec", "label": "Nội dung vụ tai nạn", "question": "Trình bày chi tiết vụ việc tai nạn giao thông (thời gian, địa điểm, diễn biến vụ việc...):", "required": True},
            {"key": "nguoi_gay_tai_nan", "label": "Người gây tai nạn", "question": "Nhập Họ tên đối tượng gây ra tai nạn giao thông (nếu không rõ, ghi 'Chưa rõ danh tính'):", "required": True},
            {"key": "hanh_vi_vi_pham", "label": "Hành vi vi phạm", "question": "Hành vi vi phạm của đối tượng (Ví dụ: đi không đúng làn đường gây tai nạn, chạy quá tốc độ...):", "required": True}
        ]
    },
    # 9. khoi_kien_gay_tai_nan_giao_thong (Mẫu 23-DS)
    {
        "id": "khoi_kien_gay_tai_nan_giao_thong",
        "category": "giao_thong",
        "file_name": "mau-don-khoi-kien-gui-gay-tai-nan-giao-thong.docx",
        "title": "Đơn khởi kiện vụ án giao thông",
        "description": "Đơn khởi kiện yêu cầu bồi thường thiệt hại ngoài hợp đồng do hành vi gây tai nạn giao thông (Mẫu số 23-DS).",
        "replacements": [
            ("……(1), ngày….. tháng …… năm…….", "{{dia_diem_ngay}}"),
            ("Kính gửi: Toà án nhân dân (2)……………………………………", "Kính gửi: Toà án nhân dân {{toa_an}}"),
            ("Người khởi kiện: (3)", "Người khởi kiện: {{nguoi_khoi_kien}}"),
            ("Địa chỉ: (4)", "Địa chỉ: {{dia_chi_khoi_kien}}"),
            ("Số điện thoại: …………………(nếu có); số fax: ………………….(nếu có)", "Số điện thoại: {{sdt_khoi_kien}} (nếu có); số fax: (nếu có)"),
            ("Địa chỉ thư điện tử: ………................................................. (nếu có)", "Địa chỉ thư điện tử: {{email_khoi_kien}} (nếu có)"),
            ("Người bị kiện: (5)", "Người bị kiện: {{nguoi_bi_kien}}"),
            ("Địa chỉ (6)", "Địa chỉ: {{dia_chi_bi_kien}}"),
            ("Người có quyền, lợi ích được bảo vệ (nếu có)(7)", "Người có quyền, lợi ích được bảo vệ (nếu có): {{nguoi_duoc_bao_ve}}"),
            ("Địa chỉ: (8)", "Địa chỉ: {{dia_chi_duoc_bao_ve}}"),
            ("Người có quyền lợi, nghĩa vụ liên quan (nếu có) (9)", "Người có quyền lợi, nghĩa vụ liên quan (nếu có): {{nguoi_lien_quan}}"),
            ("Địa chỉ: (10)", "Địa chỉ: {{dia_chi_lien_quan}}"),
            ("Yêu cầu Tòa án giải quyết những vấn đề sau đây:(11)", "Yêu cầu Tòa án giải quyết những vấn đề sau đây: {{yeu_cau_toa_an}}"),
            ("Người làm chứng (nếu có) (12)", "Người làm chứng (nếu có): {{nguoi_lam_chung}}"),
            ("Địa chỉ: (13)", "Địa chỉ: {{dia_chi_lam_chung}}"),
            ("Danh mục tài liệu, chứng kèm theo đơn khởi kiện gồm có: (14)", "Danh mục tài liệu, chứng kèm theo đơn khởi kiện gồm có: {{danh_muc_tai_lieu}}"),
            ("(Các thông tin khác mà người khởi kiện xét thấy cần thiết cho việc giải quyết vụ án) (15)", "(Các thông tin khác mà người khởi kiện xét thấy cần thiết cho việc giải quyết vụ án): {{thong_tin_khac}}"),
        ],
        "fields": [
            {"key": "dia_diem_ngay", "label": "Địa điểm và ngày làm đơn", "question": "Nhập địa điểm và ngày tháng năm làm đơn (Ví dụ: Hà Nội, ngày 18 tháng 5 năm 2026):", "required": True},
            {"key": "toa_an", "label": "Tòa án có thẩm quyền", "question": "Kính gửi Tòa án nhân dân quận/huyện/tỉnh nào?", "required": True},
            {"key": "nguoi_khoi_kien", "label": "Họ tên người khởi kiện", "question": "Nhập Họ và tên đầy đủ của người khởi kiện (bạn):", "required": True},
            {"key": "dia_chi_khoi_kien", "label": "Địa chỉ người khởi kiện", "question": "Nhập địa chỉ cư trú đầy đủ của người khởi kiện:", "required": True},
            {"key": "sdt_khoi_kien", "label": "SĐT người khởi kiện", "question": "Nhập số điện thoại của người khởi kiện:", "required": True},
            {"key": "email_khoi_kien", "label": "Email người khởi kiện", "question": "Nhập địa chỉ email của người khởi kiện (nếu có):", "required": False},
            {"key": "nguoi_bi_kien", "label": "Họ tên người bị kiện", "question": "Nhập Họ và tên đầy đủ của người gây ra tai nạn (người bị kiện):", "required": True},
            {"key": "dia_chi_bi_kien", "label": "Địa chỉ người bị kiện", "question": "Nhập địa chỉ của người bị kiện:", "required": True},
            {"key": "nguoi_duoc_bao_ve", "label": "Người được bảo vệ (nếu có)", "question": "Nhập Họ tên người có quyền lợi ích được bảo vệ (nếu có):", "required": False},
            {"key": "dia_chi_duoc_bao_ve", "label": "Địa chỉ người được bảo vệ", "question": "Địa chỉ người được bảo vệ (nếu có):", "required": False},
            {"key": "nguoi_lien_quan", "label": "Người có quyền lợi nghĩa vụ liên quan", "question": "Họ tên người có quyền lợi, nghĩa vụ liên quan (nếu có):", "required": False},
            {"key": "dia_chi_lien_quan", "label": "Địa chỉ người liên quan", "question": "Địa chỉ người có quyền lợi nghĩa vụ liên quan (nếu có):", "required": False},
            {"key": "yeu_cau_toa_an", "label": "Yêu cầu bồi thường thiệt hại", "question": "Nhập các nội dung yêu cầu Tòa án giải quyết (Ví dụ: Buộc bị đơn bồi thường chi phí sửa chữa xe, chi phí điều trị vết thương tổng cộng 50 triệu đồng...):", "required": True},
            {"key": "nguoi_lam_chung", "label": "Người làm chứng (nếu có)", "question": "Họ tên người làm chứng (nếu có):", "required": False},
            {"key": "dia_chi_lam_chung", "label": "Địa chỉ người làm chứng", "question": "Địa chỉ người làm chứng (nếu có):", "required": False},
            {"key": "danh_muc_tai_lieu", "label": "Danh mục tài liệu chứng cứ", "question": "Nhập các tài liệu, chứng cứ kèm theo đơn (Ví dụ: Biên bản tai nạn giao thông, Hóa đơn viện phí, Hóa đơn sửa xe...):", "required": True},
            {"key": "thong_tin_khac", "label": "Các thông tin khác", "question": "Các thông tin khác xét thấy cần thiết (nếu có):", "required": False}
        ]
    },
    # 10. don_thuan_tinh_ly_hon_mau (Dân sự việc)
    {
        "id": "don_thuan_tinh_ly_hon_mau",
        "category": "khac",
        "file_name": "don-thuan-tinh-ly-hon.docx",
        "title": "Đơn thuận tình ly hôn chính thức",
        "description": "Đơn yêu cầu công nhận thuận tình ly hôn, thỏa thuận nuôi con và chia tài sản (Mẫu đơn chuẩn gửi Tòa án).",
        "replacements": [
            ("(V/v: …………………………………..)(1)", "(V/v: Công nhận thuận tình ly hôn, thỏa thuận nuôi con và chia tài sản)"),
            ("Kính gửi: Tòa án nhân dân…………………(2)", "Kính gửi: Tòa án nhân dân {{toa_an}}"),
            ("Người yêu cầu giải quyết việc dân sự:(3) ...............................................................................", "Người yêu cầu giải quyết việc dân sự: {{ho_ten_vo_chong}}"),
            ("Địa chỉ:(4) ..............................................................................................................................", "Địa chỉ: {{dia_chi}}"),
            ("Số điện thoại (nếu có): …………………………..; Fax (nếu có):................................................", "Số điện thoại (nếu có): {{sdt}}"),
            ("Địa chỉ thư điện tử (nếu có): ..................................................................................................", "Địa chỉ thư điện tử (nếu có): {{email}}"),
            ("Tôi (chúng tôi) xin trình bày với Tòa án nhân dân(5) ................................................................", "Tôi (chúng tôi) xin trình bày với Tòa án nhân dân {{toa_an}}"),
            ("- Những vấn đề yêu cầu Tòa án giải quyết:(6) .........................................................................", "- Những vấn đề yêu cầu Tòa án giải quyết: {{yeu_cau_toa_an}}"),
            ("- Lý do, mục đích, căn cứ của việc yêu cầu Tòa án giải quyết đối với những vấn đề nêu trên:(7)", "- Lý do, mục đích, căn cứ của việc yêu cầu Tòa án giải quyết đối với những vấn đề nêu trên: {{ly_do_yeu_cau}}"),
            ("- Tên và địa chỉ của những người có liên quan đến những vấn đề yêu cầu Tòa án giải quyết:(8)  ", "- Tên và địa chỉ của những người có liên quan đến những vấn đề yêu cầu Tòa án giải quyết: {{nguoi_lien_quan}}"),
            ("- Các thông tin khác (nếu có):(9).............................................................................................", "- Các thông tin khác (nếu có): {{thong_tin_khac}}"),
            ("Tài liệu, chứng cứ kèm theo đơn yêu cầu:(10)", "Tài liệu, chứng cứ kèm theo đơn yêu cầu: {{chung_cu_kem_theo}}"),
            ("1. .........................................................................................................................................", "1. {{chung_cu_1}}"),
            ("2. .........................................................................................................................................", "2. {{chung_cu_2}}"),
            ("3. .........................................................................................................................................", "3. {{chung_cu_3}}"),
            ("……, ngày…. tháng…. năm……. (11)", "{{ngay_lam_don}}")
        ],
        "fields": [
            {"key": "toa_an", "label": "Tòa án có thẩm quyền", "question": "Kính gửi Tòa án nhân dân quận/huyện/tỉnh nào giải quyết việc ly hôn?", "required": True},
            {"key": "ho_ten_vo_chong", "label": "Họ tên vợ và chồng", "question": "Nhập Họ tên vợ và Họ tên chồng (Ngăn cách bởi dấu phẩy, Ví dụ: Nguyễn Văn A, Trần Thị B):", "required": True},
            {"key": "dia_chi", "label": "Địa chỉ cư trú của vợ chồng", "question": "Nhập địa chỉ cư trú đầy đủ hiện tại của hai vợ chồng:", "required": True},
            {"key": "sdt", "label": "Số điện thoại liên lạc", "question": "Nhập số điện thoại liên lạc của hai vợ chồng:", "required": True},
            {"key": "email", "label": "Email liên lạc", "question": "Nhập địa chỉ thư điện tử của hai vợ chồng (nếu có):", "required": False},
            {"key": "yeu_cau_toa_an", "label": "Các vấn đề yêu cầu giải quyết", "question": "Nhập các vấn đề yêu cầu Tòa án giải quyết (Ví dụ: Công nhận thuận tình ly hôn; Thỏa thuận quyền nuôi con; Thỏa thuận phân chia tài sản chung...):", "required": True},
            {"key": "ly_do_yeu_cau", "label": "Lý do và căn cứ ly hôn", "question": "Trình bày tóm tắt lý do ly hôn (Ví dụ: Vợ chồng bất đồng quan điểm, mục đích hôn nhân không đạt được...):", "required": True},
            {"key": "nguoi_lien_quan", "label": "Người liên quan (nếu có)", "question": "Họ tên và địa chỉ của những người có liên quan (Ví dụ: các con chung...):", "required": False},
            {"key": "thong_tin_khac", "label": "Các thông tin khác", "question": "Các thông tin khác cần trình bày (nếu có):", "required": False},
            {"key": "chung_cu_kem_theo", "label": "Tài liệu chứng cứ kèm theo", "question": "Mô tả chung các chứng cứ tài liệu kèm theo (Ví dụ: Đăng ký kết hôn bản gốc, CCCD bản sao...):", "required": True},
            {"key": "chung_cu_1", "label": "Chứng cứ kèm theo 1", "question": "Nhập tài liệu chứng cứ kèm theo số 1 (Ví dụ: Bản chính Giấy chứng nhận đăng ký kết hôn):", "required": True},
            {"key": "chung_cu_2", "label": "Chứng cứ kèm theo 2", "question": "Nhập tài liệu chứng cứ kèm theo số 2 (Ví dụ: Bản sao CCCD của vợ và chồng):", "required": True},
            {"key": "chung_cu_3", "label": "Chứng cứ kèm theo 3", "question": "Nhập tài liệu chứng cứ kèm theo số 3 (Ví dụ: Bản sao Giấy khai sinh của các con):", "required": False},
            {"key": "ngay_lam_don", "label": "Ngày làm đơn", "question": "Nhập địa điểm và ngày làm đơn (Ví dụ: Hà Nội, ngày 18 tháng 5 năm 2026):", "required": True}
        ]
    },
    # 11. don_xin_hoc_them
    {
        "id": "don_xin_hoc_them",
        "category": "khac",
        "file_name": "don-xin-hoc-them-mon-hoc (1).docx",
        "title": "Đơn xin học lớp bồi dưỡng kiến thức",
        "description": "Đơn của học sinh xin đăng ký tham gia lớp học bồi dưỡng kiến thức/học thêm môn học do Nhà trường tổ chức.",
        "replacements": [
            ("Kính gửi: Ban Giám hiệu Trường ..............", "Kính gửi: Ban Giám hiệu Trường {{truong}}"),
            ("Tên em là: ...................................................................................................................", "Tên em là: {{ten_hoc_sinh}}"),
            ("Học sinh lớp .............. Trường ..........................................................................", "Học sinh lớp {{lop}} Trường {{truong_hien_tai}}"),
            ("trong đó có môn ………… cho học sinh khối ......", "trong đó có môn {{mon_hoc}} cho học sinh khối {{khoi}}"),
            ("..........., ngày...tháng....năm...", "{{ngay_lam_don}}")
        ],
        "fields": [
            {"key": "truong", "label": "Tên trường", "question": "Nhập tên trường học nơi tổ chức lớp bồi dưỡng:", "required": True},
            {"key": "ten_hoc_sinh", "label": "Tên học sinh", "question": "Nhập Họ và tên đầy đủ của học sinh:", "required": True},
            {"key": "lop", "label": "Lớp học hiện tại", "question": "Nhập tên lớp học hiện tại (Ví dụ: 10A1):", "required": True},
            {"key": "truong_hien_tai", "label": "Trường học hiện tại", "question": "Nhập tên trường học hiện tại của học sinh:", "required": True},
            {"key": "mon_hoc", "label": "Môn học đăng ký", "question": "Nhập tên môn học xin học bồi dưỡng (Ví dụ: Tiếng Anh):", "required": True},
            {"key": "khoi", "label": "Khối học", "question": "Nhập khối học hiện tại (Ví dụ: 10):", "required": True},
            {"key": "ngay_lam_don", "label": "Ngày làm đơn", "question": "Nhập địa điểm và ngày làm đơn (Ví dụ: Hà Nội, ngày 18 tháng 5 năm 2026):", "required": True}
        ]
    },
    # 12. don_xin_xac_nhan_gia_dinh_kho_khan
    {
        "id": "don_xin_xac_nhan_gia_dinh_kho_khan",
        "category": "khac",
        "file_name": "mau-don-xin-xac-nhan-gia-dinh-kho-khan.docx",
        "title": "Đơn xin xác nhận hoàn cảnh khó khăn",
        "description": "Đơn xin UBND cấp xã/phường xác nhận hoàn cảnh kinh tế gia đình khó khăn để phục vụ xin học bổng, miễn giảm học phí.",
        "replacements": [
            ("Kính gửi: UBND xã, (phường):..................................................", "Kính gửi: UBND xã, (phường): {{ubnd_xa_phuong}}"),
            ("1. Tôi tên là:....................................... Sinh ngày:..........................................", "1. Tôi tên là: {{ho_ten}} Sinh ngày: {{ngay_sinh}}"),
            ("2. Quê quán:.................................................................................................", "2. Quê quán: {{que_quan}}"),
            ("3. Địa chỉ (tạm trú):………………………………………………………………", "3. Địa chỉ (tạm trú): {{dia_chi_tam_tru}}"),
            ("4. Nghề nghiệp: ...........................................................................................", "4. Nghề nghiệp: {{nghe_nghiep}}"),
            ("Xin được xác nhận gia đình tôi hiện đang có hoàn cảnh như sau:", "Xin được xác nhận gia đình tôi hiện đang có hoàn cảnh như sau: {{hoan_canh_chi_tiet}}"),
            ("Cha:................................................ tuổi, hiện ở tại:......................", "Cha: {{tuoi_cha}} tuổi, hiện ở tại: {{dia_chi_cha}}"),
            ("Nghề nghiệp: ................................................................................", "Nghề nghiệp: {{nghe_nghiep_cha}}"),
            ("Mẹ:.................................................. tuổi, hiện ở tại:......................", "Mẹ: {{tuoi_me}} tuổi, hiện ở tại: {{dia_chi_me}}"),
            ("Nghề nghiệp:..................................................................................", "Nghề nghiệp: {{nghe_nghiep_me}}"),
            ("Khác (cha mẹ ly thân, ly hôn…):.....................................................", "Khác (cha mẹ ly thân, ly hôn…): {{hoan_canh_bo_me_khac}}"),
            ("6. Gia đình có……anh chị em (kể cả tôi). Người lớn nhất….............….tuổi, Người nhỏ nhất...........……tuổi.", "6. Gia đình có {{so_anh_chi_em}} anh chị em (kể cả tôi). Người lớn nhất {{tuoi_lon_nhat}} tuổi, Người nhỏ nhất {{tuoi_nho_nhat}} tuổi."),
            ("Số người còn đang đi học: Cấp 1:….......……, Cấp 2:…..….....…, Cấp 3:…………, Đại học……………", "Số người còn đang đi học: Cấp 1: {{di_hoc_cap_1}}, Cấp 2: {{di_hoc_cap_2}}, Cấp 3: {{di_hoc_cap_3}}, Đại học: {{di_hoc_dai_hoc}}"),
            ("7. Nhà tôi có ......…m2 đất trồng (nuôi)….. …....…………………………..", "7. Nhà tôi có {{dien_tich_dat}} m2 đất trồng (nuôi)"),
            ("8. Gia đình có buôn bán nhỏ (nghề nghiệp khác)….……………...", "8. Gia đình có buôn bán nhỏ (nghề nghiệp khác): {{gia_dinh_buon_ban}}"),
            ("9. Thu nhập bình quân của gia đình……………………….đồng/tháng", "9. Thu nhập bình quân của gia đình: {{thu_nhap_gia_dinh}} đồng/tháng"),
            ("10. Bản thân tôi: Thu nhập…………………….đồng/tháng (nếu có)", "10. Bản thân tôi: Thu nhập: {{thu_nhap_ban_than}} đồng/tháng (nếu có)"),
            ("11. Lý do xin xác nhận: (VD: bổ sung hồ sơ xin học bổng, vay vốn….)", "11. Lý do xin xác nhận: {{ly_do_xin_xac_nhan}}"),
            ("…, ngày…tháng……năm 20....", "{{ngay_lam_don}}")
        ],
        "fields": [
            {"key": "ubnd_xa_phuong", "label": "UBND xã/phường", "question": "Kính gửi Ủy ban nhân dân xã/phường/thị trấn nào?", "required": True},
            {"key": "ho_ten", "label": "Họ và tên của bạn", "question": "Nhập Họ và tên đầy đủ của bạn:", "required": True},
            {"key": "ngay_sinh", "label": "Ngày sinh", "question": "Nhập ngày tháng năm sinh của bạn (dd/mm/yyyy):", "required": True},
            {"key": "que_quan", "label": "Quê quán", "question": "Nhập quê quán của bạn (Tỉnh/Thành phố):", "required": True},
            {"key": "dia_chi_tam_tru", "label": "Địa chỉ cư trú", "question": "Nhập địa chỉ cư trú hoặc địa chỉ tạm trú hiện tại của bạn:", "required": True},
            {"key": "nghe_nghiep", "label": "Nghề nghiệp của bạn", "question": "Nhập nghề nghiệp hiện tại của bạn (Ví dụ: Sinh viên):", "required": True},
            {"key": "hoan_canh_chi_tiet", "label": "Chi tiết hoàn cảnh gia đình", "question": "Mô tả chi tiết về hoàn cảnh khó khăn của gia đình bạn (Ví dụ: Thu nhập thấp, cha mẹ đau ốm thường xuyên...):", "required": True},
            {"key": "tuoi_cha", "label": "Tuổi của cha", "question": "Nhập tuổi của cha bạn (ghi năm mất và lý do nếu đã mất):", "required": True},
            {"key": "dia_chi_cha", "label": "Địa chỉ hiện tại của cha", "question": "Địa chỉ cư trú hiện tại của cha bạn:", "required": True},
            {"key": "nghe_nghiep_cha", "label": "Nghề nghiệp của cha", "question": "Nghề nghiệp của cha bạn:", "required": True},
            {"key": "tuoi_me", "label": "Tuổi của mẹ", "question": "Nhập tuổi của mẹ bạn (ghi năm mất và lý do nếu đã mất):", "required": True},
            {"key": "dia_chi_me", "label": "Địa chỉ hiện tại của mẹ", "question": "Địa chỉ cư trú hiện tại của mẹ bạn:", "required": True},
            {"key": "nghe_nghiep_me", "label": "Nghề nghiệp của mẹ", "question": "Nghề nghiệp của mẹ bạn:", "required": True},
            {"key": "hoan_canh_bo_me_khac", "label": "Hoàn cảnh cha mẹ khác", "question": "Trình bày hoàn cảnh khác của cha mẹ (Ví dụ: ly thân, ly hôn, bệnh tật, thương binh...):", "required": False},
            {"key": "so_anh_chi_em", "label": "Số anh chị em", "question": "Tổng số anh chị em trong gia đình bạn (kể cả bạn):", "required": True},
            {"key": "tuoi_lon_nhat", "label": "Tuổi người lớn nhất", "question": "Tuổi của người anh/chị lớn nhất trong gia đình:", "required": True},
            {"key": "tuoi_nho_nhat", "label": "Tuổi người nhỏ nhất", "question": "Tuổi của người em nhỏ nhất trong gia đình:", "required": True},
            {"key": "di_hoc_cap_1", "label": "Số người đi học Cấp 1", "question": "Số lượng người đang đi học Cấp 1 trong gia đình:", "required": True},
            {"key": "di_hoc_cap_2", "label": "Số người đi học Cấp 2", "question": "Số lượng người đang đi học Cấp 2 trong gia đình:", "required": True},
            {"key": "di_hoc_cap_3", "label": "Số người đi học Cấp 3", "question": "Số lượng người đang đi học Cấp 3 trong gia đình:", "required": True},
            {"key": "di_hoc_dai_hoc", "label": "Số người học Đại học", "question": "Số lượng người đang học Đại học trong gia đình:", "required": True},
            {"key": "dien_tich_dat", "label": "Diện tích đất trồng", "question": "Diện tích đất trồng trọt hoặc chăn nuôi của gia đình (m2) (nếu không có ghi 0):", "required": True},
            {"key": "gia_dinh_buon_ban", "label": "Gia đình có kinh doanh", "question": "Gia đình có buôn bán nhỏ hay làm nghề phụ nào khác không?", "required": True},
            {"key": "thu_nhap_gia_dinh", "label": "Thu nhập của gia đình", "question": "Nhập thu nhập bình quân của hộ gia đình (đồng/tháng):", "required": True},
            {"key": "thu_nhap_ban_than", "label": "Thu nhập cá nhân", "question": "Thu nhập của riêng bạn (đồng/tháng) (nếu có, không có ghi 0):", "required": False},
            {"key": "ly_do_xin_xac_nhan", "label": "Lý do xin xác nhận", "question": "Mục đích xin xác nhận này để làm gì? (Ví dụ: Để nộp hồ sơ xét miễn giảm học phí):", "required": True},
            {"key": "ngay_lam_don", "label": "Ngày làm đơn", "question": "Nhập địa điểm và ngày làm đơn (Ví dụ: Hà Nội, ngày 18 tháng 5 năm 2026):", "required": True}
        ]
    },
    # 13. don_ly_hon_don_phuong_mau (Mẫu 23-DS)
    {
        "id": "don_ly_hon_don_phuong_mau",
        "category": "khac",
        "file_name": "Đơn ly hôn đơn phương.docx",
        "title": "Đơn ly hôn đơn phương chính thức",
        "description": "Đơn khởi kiện vụ án ly hôn đơn phương (Mẫu số 23-DS) gửi Tòa án nhân dân giải quyết.",
        "replacements": [
            ("……(1), ngày….. tháng …… năm…….", "{{dia_diem_ngay}}"),
            ("Kính gửi: Toà án nhân dân (2)……………………………………", "Kính gửi: Toà án nhân dân {{toa_an}}"),
            ("Người khởi kiện: (3).........................................................................................................................", "Người khởi kiện: {{nguoi_khoi_kien}}"),
            ("Địa chỉ: (4) .......................................................................................................................................", "Địa chỉ: {{dia_chi_khoi_kien}}"),
            ("Số điện thoại: …………………(nếu có); số fax: ………………….(nếu có)", "Số điện thoại: {{sdt_khoi_kien}} (nếu có); số fax: (nếu có)"),
            ("Địa chỉ thư điện tử: ………................................................. (nếu có)", "Địa chỉ thư điện tử: {{email_khoi_kien}} (nếu có)"),
            ("Người bị kiện: (5).............................................................................................................................", "Người bị kiện: {{nguoi_bi_kien}}"),
            ("Địa chỉ (6) ........................................................................................................................................", "Địa chỉ: {{dia_chi_bi_kien}}"),
            ("Người có quyền, lợi ích được bảo vệ (nếu có)(7).............................................................................", "Người có quyền, lợi ích được bảo vệ (nếu có): {{nguoi_duoc_bao_ve}}"),
            ("Địa chỉ: (8)........................................................................................................................................", "Địa chỉ: {{dia_chi_duoc_bao_ve}}"),
            ("Người có quyền lợi, nghĩa vụ liên quan (nếu có) (9)........................................................................", "Người có quyền lợi, nghĩa vụ liên quan (nếu có): {{nguoi_lien_quan}}"),
            ("Địa chỉ: (10) ......................................................................................................................................", "Địa chỉ: {{dia_chi_lien_quan}}"),
            ("Yêu cầu Tòa án giải quyết những vấn đề sau đây:(11)......................................................................", "Yêu cầu Tòa án giải quyết những vấn đề sau đây: {{yeu_cau_toa_an}}"),
            ("Người làm chứng (nếu có) (12)...........................................................................................................", "Người làm chứng (nếu có): {{nguoi_lam_chung}}"),
            ("Địa chỉ: (13) ......................................................................................................................................", "Địa chỉ: {{dia_chi_lam_chung}}"),
            ("Danh mục tài liệu, chứng kèm theo đơn khởi kiện gồm có: (14).......................................................", "Danh mục tài liệu, chứng kèm theo đơn khởi kiện gồm có: {{danh_muc_tai_lieu}}"),
            ("1.....................................................................................................................................................", "1. {{danh_muc_1}}"),
            ("2.....................................................................................................................................................", "2. {{danh_muc_2}}"),
            ("(Các thông tin khác mà người khởi kiện xét thấy cần thiết cho việc giải quyết vụ án) (15) .....................", "(Các thông tin khác mà người khởi kiện xét thấy cần thiết cho việc giải quyết vụ án): {{thong_tin_khac}}"),
        ],
        "fields": [
            {"key": "dia_diem_ngay", "label": "Địa điểm và ngày làm đơn", "question": "Nhập địa điểm và ngày tháng năm làm đơn (Ví dụ: Hà Nội, ngày 18 tháng 5 năm 2026):", "required": True},
            {"key": "toa_an", "label": "Tòa án có thẩm quyền", "question": "Kính gửi Tòa án nhân dân quận/huyện/tỉnh nào?", "required": True},
            {"key": "nguoi_khoi_kien", "label": "Họ tên người khởi kiện (vợ/chồng)", "question": "Nhập Họ và tên đầy đủ của người khởi kiện (bạn):", "required": True},
            {"key": "dia_chi_khoi_kien", "label": "Địa chỉ người khởi kiện", "question": "Nhập địa chỉ cư trú đầy đủ hiện tại của bạn:", "required": True},
            {"key": "sdt_khoi_kien", "label": "SĐT người khởi kiện", "question": "Nhập số điện thoại của bạn:", "required": True},
            {"key": "email_khoi_kien", "label": "Email người khởi kiện", "question": "Nhập địa chỉ email của bạn (nếu có):", "required": False},
            {"key": "nguoi_bi_kien", "label": "Họ tên người bị kiện (chồng/vợ)", "question": "Nhập Họ và tên đầy đủ của người bị khởi kiện (vợ hoặc chồng của bạn):", "required": True},
            {"key": "dia_chi_bi_kien", "label": "Địa chỉ người bị kiện", "question": "Nhập địa chỉ cư trú đầy đủ hiện tại của người bị kiện:", "required": True},
            {"key": "nguoi_duoc_bao_ve", "label": "Người được bảo vệ (nếu có)", "question": "Nhập Họ tên người có quyền lợi ích được bảo vệ (nếu có):", "required": False},
            {"key": "dia_chi_duoc_bao_ve", "label": "Địa chỉ người được bảo vệ", "question": "Địa chỉ người được bảo vệ (nếu có):", "required": False},
            {"key": "nguoi_lien_quan", "label": "Người có quyền lợi nghĩa vụ liên quan", "question": "Họ tên người có quyền lợi, nghĩa vụ liên quan (nếu có):", "required": False},
            {"key": "dia_chi_lien_quan", "label": "Địa chỉ người liên quan", "question": "Địa chỉ người có quyền lợi nghĩa vụ liên quan (nếu có):", "required": False},
            {"key": "yeu_cau_toa_an", "label": "Yêu cầu giải quyết ly hôn", "question": "Nhập các nội dung yêu cầu Tòa án giải quyết (Ví dụ: Yêu cầu giải quyết ly hôn đơn phương với ông/bà B, yêu cầu trực tiếp nuôi dưỡng con chung và không yêu cầu chia tài sản...):", "required": True},
            {"key": "nguoi_lam_chung", "label": "Người làm chứng (nếu có)", "question": "Họ tên người làm chứng (nếu có):", "required": False},
            {"key": "dia_chi_lam_chung", "label": "Địa chỉ người làm chứng", "question": "Địa chỉ người làm chứng (nếu có):", "required": False},
            {"key": "danh_muc_tai_lieu", "label": "Danh mục tài liệu chứng cứ", "question": "Nhập các tài liệu, chứng cứ kèm theo đơn (Ví dụ: Đăng ký kết hôn bản chính, Giấy khai sinh các con bản sao...):", "required": True},
            {"key": "danh_muc_1", "label": "Chứng cứ 1", "question": "Tên tài liệu chứng cứ kèm theo số 1 (Ví dụ: Bản chính Giấy chứng nhận kết hôn):", "required": True},
            {"key": "danh_muc_2", "label": "Chứng cứ 2", "question": "Tên tài liệu chứng cứ kèm theo số 2 (Ví dụ: Bản sao Giấy khai sinh con chung):", "required": True},
            {"key": "thong_tin_khac", "label": "Các thông tin khác", "question": "Các thông tin khác xét thấy cần thiết (nếu có):", "required": False}
        ]
    }
]

# Run processing
for spec in templates_specs:
    print(f"Processing spec: {spec['id']}...")
    src = os.path.join(BASE_MAUDON_DIR, spec['category'], spec['file_name'])
    if not os.path.isfile(src):
        # Fallback to category folder checking or check if it exists in 'khac'
        if spec['category'] == 'khac' or spec['category'] == 'dan_su':
            # Check other folders
            for possible_cat in ['lao_dong', 'dan_su', 'giao_thong', 'khac']:
                temp_src = os.path.join(BASE_MAUDON_DIR, possible_cat, spec['file_name'])
                if os.path.isfile(temp_src):
                    src = temp_src
                    break
        
        if not os.path.isfile(src):
            print(f"  [ERROR] File not found: {src}")
            continue
            
    dst_docx_name = f"{spec['id']}_template.docx"
    dst_docx_path = os.path.join(TARGET_TEMPLATE_DIR, dst_docx_name)
    
    # Process replacements in DOCX
    replace_docx_placeholders(src, dst_docx_path, spec['replacements'])
    
    # Generate JSON template
    json_path = os.path.join(TARGET_TEMPLATE_DIR, f"{spec['id']}.json")
    json_data = {
        "id": spec['id'],
        "title": spec['title'],
        "description": spec['description'],
        "docx_template": dst_docx_name,
        "fields": spec['fields'],
        "output_template": "Bạn đã hoàn thành mẫu đơn: " + spec['title'] + ".\n\nVui lòng tải xuống file Word để chỉnh sửa chi tiết."
    }
    
    with open(json_path, 'w', encoding='utf-8') as jf:
        json.dump(json_data, jf, ensure_ascii=False, indent=4)
    print(f"  Successfully created JSON template: {json_path}")

print("=== ALL PROCESSES COMPLETED SUCCESSFULLY ===")
